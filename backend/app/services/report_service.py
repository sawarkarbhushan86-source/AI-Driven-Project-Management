import os
import json
from datetime import datetime, timezone
from typing import Dict, Any, Optional
from sqlalchemy.orm import Session
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models.project import Project
from app.models.task import Task
from app.models.task_update import TaskUpdate
from app.models.project_risk import ProjectRisk
from app.models.report import Report
from app.services.genai_service import genai_service

class ReportExportService:
    def __init__(self):
        self.reports_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "reports")
        os.makedirs(self.reports_dir, exist_ok=True)

    def generate_and_save_report(
        self,
        project_id: int,
        report_type: str,
        user_id: int,
        custom_notes: str,
        db: Session
    ) -> Report:
        project = db.query(Project).filter(Project.id == project_id).first()
        if not project:
            raise ValueError("Project not found")

        tasks = db.query(Task).filter(Task.project_id == project_id).all()
        risks = db.query(ProjectRisk).filter(ProjectRisk.project_id == project_id).all()
        
        completed = sum(1 for t in tasks if t.status == "COMPLETED")
        blocked = sum(1 for t in tasks if t.status == "BLOCKED")
        in_progress = sum(1 for t in tasks if t.status == "IN_PROGRESS")
        todo = sum(1 for t in tasks if t.status == "TODO")
        total = max(1, len(tasks))

        completion_pct = round((completed / total) * 100, 1)

        summary_text = (
            f"AI Executive Status Summary for {project.name}:\n"
            f"Overall health score is {project.health_score}/100. "
            f"The delay prediction engine forecasts {project.predicted_delay_days} days of schedule variance ({project.delay_risk_level} risk). "
            f"Completion is currently at {completion_pct}% with {completed} finished tasks, "
            f"{in_progress} in progress, {blocked} blocked, and {todo} queued. "
            f"{custom_notes if custom_notes else ''}"
        )

        recommendations_text = (
            "1. Focus developer resources on clearing active blockers.\n"
            "2. Conduct mid-sprint scope verification for high-risk deliverables.\n"
            "3. Maintain code review turnaround within 24 hours."
        )

        metrics = {
            "total_tasks": len(tasks),
            "completed_tasks": completed,
            "blocked_tasks": blocked,
            "completion_percentage": completion_pct,
            "health_score": float(project.health_score),
            "predicted_delay_days": float(project.predicted_delay_days),
            "risk_level": project.delay_risk_level
        }

        # Timestamped filenames
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        pdf_name = f"report_proj_{project_id}_{report_type.lower()}_{timestamp_str}.pdf"
        excel_name = f"report_proj_{project_id}_{report_type.lower()}_{timestamp_str}.xlsx"

        pdf_path = os.path.join(self.reports_dir, pdf_name)
        excel_path = os.path.join(self.reports_dir, excel_name)

        # Generate files
        self._build_pdf(project, tasks, risks, metrics, summary_text, pdf_path)
        self._build_excel(project, tasks, risks, metrics, excel_path)

        report = Report(
            project_id=project.id,
            generated_by_id=user_id,
            report_type=report_type,
            title=f"{report_type.capitalize()} Progress & AI Health Audit - {project.name}",
            summary=summary_text,
            metrics_json=metrics,
            recommendations=recommendations_text,
            pdf_filename=pdf_name,
            excel_filename=excel_name
        )
        db.add(report)
        db.commit()
        db.refresh(report)
        return report

    def _build_pdf(self, project, tasks, risks, metrics, summary, output_path):
        doc = SimpleDocTemplate(output_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=22,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=12
        )
        subtitle_style = ParagraphStyle(
            'SubtitleStyle',
            parent=styles['Normal'],
            fontSize=11,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=18
        )
        section_style = ParagraphStyle(
            'SectionStyle',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor("#0F172A"),
            spaceBefore=14,
            spaceAfter=8
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor("#334155"),
            leading=14
        )

        story.append(Paragraph(f"Project Health & Progress Audit: {project.name}", title_style))
        story.append(Paragraph(f"Generated on {datetime.now().strftime('%B %d, %Y')} | Platform: AI-Driven Project Manager", subtitle_style))

        # Metrics Table
        metrics_data = [
            ["Health Score", f"{metrics['health_score']} / 100", "Completion %", f"{metrics['completion_percentage']}%"],
            ["Predicted Delay", f"+{metrics['predicted_delay_days']} Days", "Risk Level", metrics['risk_level']],
            ["Total Tasks", str(metrics['total_tasks']), "Completed / Blocked", f"{metrics['completed_tasks']} / {metrics['blocked_tasks']}"]
        ]
        t = Table(metrics_data, colWidths=[130, 130, 130, 130])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
            ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor("#1E293B")),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        story.append(t)
        story.append(Spacer(1, 14))

        # Executive Summary
        story.append(Paragraph("AI Executive Briefing", section_style))
        story.append(Paragraph(summary, body_style))
        story.append(Spacer(1, 14))

        # Tasks Overview
        story.append(Paragraph("Recent Sprint Tasks", section_style))
        task_rows = [["ID", "Title", "Status", "Priority", "Progress"]]
        for tsk in tasks[:8]:
            task_rows.append([
                str(tsk.id),
                tsk.title[:36] + "..." if len(tsk.title) > 36 else tsk.title,
                tsk.status,
                tsk.priority,
                f"{tsk.progress_percentage}%"
            ])
        task_table = Table(task_rows, colWidths=[35, 260, 85, 70, 70])
        task_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0284C7")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('TOPPADDING', (0,0), (-1,-1), 5),
        ]))
        story.append(task_table)

        doc.build(story)

    def _build_excel(self, project, tasks, risks, metrics, output_path):
        wb = openpyxl.Workbook()

        # Sheet 1: Project Overview
        ws_overview = wb.active
        ws_overview.title = "Project Overview"

        header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

        ws_overview["A1"] = "Metric"
        ws_overview["B1"] = "Value"
        ws_overview["A1"].font = header_font
        ws_overview["A1"].fill = header_fill
        ws_overview["B1"].font = header_font
        ws_overview["B1"].fill = header_fill

        overview_items = [
            ("Project Name", project.name),
            ("Status", project.status),
            ("Health Score", metrics["health_score"]),
            ("Predicted Delay (Days)", metrics["predicted_delay_days"]),
            ("Risk Level", metrics["risk_level"]),
            ("Total Tasks", metrics["total_tasks"]),
            ("Completed Tasks", metrics["completed_tasks"]),
            ("Blocked Tasks", metrics["blocked_tasks"]),
            ("Completion Percentage", f"{metrics['completion_percentage']}%"),
            ("Budget", float(project.budget or 0)),
            ("Start Date", str(project.start_date)),
            ("Deadline", str(project.deadline))
        ]

        for r_idx, (k, v) in enumerate(overview_items, start=2):
            ws_overview.cell(row=r_idx, column=1, value=k).font = Font(bold=True)
            ws_overview.cell(row=r_idx, column=2, value=v)

        # Sheet 2: Tasks List
        ws_tasks = wb.create_sheet(title="Tasks")
        task_headers = ["ID", "Title", "Status", "Priority", "Assigned To ID", "Progress %", "Est Hours", "Actual Hours", "At Risk"]
        for col_idx, h in enumerate(task_headers, start=1):
            cell = ws_tasks.cell(row=1, column=col_idx, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")

        for r_idx, t in enumerate(tasks, start=2):
            ws_tasks.cell(row=r_idx, column=1, value=t.id)
            ws_tasks.cell(row=r_idx, column=2, value=t.title)
            ws_tasks.cell(row=r_idx, column=3, value=t.status)
            ws_tasks.cell(row=r_idx, column=4, value=t.priority)
            ws_tasks.cell(row=r_idx, column=5, value=t.assigned_to_id)
            ws_tasks.cell(row=r_idx, column=6, value=t.progress_percentage)
            ws_tasks.cell(row=r_idx, column=7, value=float(t.estimated_hours or 0))
            ws_tasks.cell(row=r_idx, column=8, value=float(t.actual_hours or 0))
            ws_tasks.cell(row=r_idx, column=9, value="YES" if t.is_at_risk else "NO")

        # Auto column width
        for ws in [ws_overview, ws_tasks]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(output_path)

report_export_service = ReportExportService()
